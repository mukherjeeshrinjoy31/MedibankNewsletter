import re
import io
import os
import json
import time
import argparse
import feedparser
import requests
import openpyxl
import boto3
from bs4 import BeautifulSoup
from datetime import datetime, timezone
from collections import defaultdict

# ---- Config -----------------------------------------------------------------
SOURCE  = "nib"
DATASET = "competitor_offers"
TIER    = "phi_industry"
BUCKET  = "p000268ds-comp-offers"
HEADERS = {"User-Agent": "Mozilla/5.0"}
SHEET   = "comp_offer.xlsx"

OFFER_PAGES = [
    {"name": "NIB Offers",              "url": "https://www.nib.com.au/health-insurance/offers"},
    {"name": "NIB Homepage",            "url": "https://www.nib.com.au/"},
    {"name": "NIB Health Insurance",    "url": "https://www.nib.com.au/health-insurance"},
    {"name": "NIB Hospital",            "url": "https://www.nib.com.au/health-insurance/hospital"},
    {"name": "NIB Extras",              "url": "https://www.nib.com.au/health-insurance/extras"},
]

OFFER_KEYWORDS = ["weeks free", "week free", "waiting period", "waiver", "promo",
                  "offer", "bonus", "discount", "join by", "ends", "new members",
                  "12wfreeww", "extras8wf", "skip the", "month wait"]

# ---- Regex Patterns ---------------------------------------------------------
WEEKS_PAT    = re.compile(r'(?:up to\s+)?(\d+(?:\+\d+)?)\s*weeks?\s*free', re.I)
WAITING_PAT  = re.compile(r'(\d+)\s*(?:and|&)\s*(\d+)\s*month.*?(?:wait|waiv)', re.I)
WAITING_PAT2 = re.compile(r'(\d+)\s*month.*?(?:wait|waiv)', re.I)
ENDDATE_PAT  = re.compile(r'(?:ends?|until|by|join by|choose.*?by|by\s+\d{1,2})\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4}|\d{1,2}\s+[A-Za-z]+\s*\d{2,4}|[A-Za-z]+\s+\d{1,2},?\s+\d{4})', re.I)
ENDDATE_PAT2 = re.compile(r'(\d{1,2}\s+[A-Za-z]{3,}\s+\d{4})', re.I)
ENDDATE_PAT3 = re.compile(r'(?:ends?|until|by|join by|choose.*?online by)\s+(\d{1,2}\s+[A-Za-z]+)(?:\s|\.)', re.I)
PROMO_PAT    = re.compile(r'(?:promo(?:tion)?\s*code|use\s+(?:promo\s+)?code|enter\s+code|code)[:\s]+([A-Z0-9]{6,})', re.I)

EXCEL_COL_MAP = {
    "weeks_free":    "Offer : Weeks Free",
    "waiting_waive": "Offer : Waiting period waive",
    "other":         "Offer : Other",
    "end_date":      "Offer : End date",
}

COVER_TYPE_KEYWORDS = {
    "hospital + extras": ["hospital and extras", "hospital & extras", "combined",
                          "hospital + extras", "hospital and extra", "12wfreeww",
                          "choose a combined"],
    "hospital only":     ["hospital only", "hospital cover", "hospital-only",
                          "standalone hospital", "hospital product",
                          "choose a hospital only", "hospital only cover"],
    "extras only":       ["extras only", "extras cover", "extras-only",
                          "standalone extras", "extras8wf", "extras policy",
                          "eligible extras", "extras cover online"],
}

# ---- Scraping ---------------------------------------------------------------
def fetch_page(url: str) -> BeautifulSoup | None:
    try:
        r = requests.get(url, headers=HEADERS, timeout=30)
        if r.status_code == 200:
            return BeautifulSoup(r.text, "html.parser")
        print(f"  ✗ {url}: {r.status_code}")
    except Exception as e:
        print(f"  ✗ {url}: {e}")
    return None


def extract_offer_blocks(soup: BeautifulSoup) -> list[str]:
    found = []
    for tag in soup.find_all(["p", "h1", "h2", "h3", "h4", "li", "span", "div"]):
        text = tag.get_text(separator=" ", strip=True)
        if any(kw in text.lower() for kw in OFFER_KEYWORDS):
            if 20 < len(text) < 600 and text not in found:
                found.append(text)
    return found[:30]


def scrape_rss() -> list[str]:
    print("  Fetching NIB offers via Google News RSS...")
    rss_url = "https://news.google.com/rss/search?q=NIB+health+insurance+offer+weeks+free+Australia&hl=en-AU&gl=AU&ceid=AU:en"
    feed = feedparser.parse(rss_url)
    entries = []
    for entry in feed.entries[:10]:
        entries.append(f"{entry.title} | {entry.link} | {entry.get('published', '')}")
    print(f"  ✓ Found {len(entries)} RSS articles")
    return entries


def scrape_all_pages() -> list[str]:
    print("--- NIB Scraper ---")
    all_blocks = []

    for page in OFFER_PAGES:
        print(f"  Fetching {page['name']}...")
        soup = fetch_page(page["url"])
        if soup:
            blocks = extract_offer_blocks(soup)
            all_blocks.extend(blocks)
            print(f"  ✓ {page['name']}: {len(blocks)} offer blocks found")
        else:
            print(f"  ✗ {page['name']}: failed")
        time.sleep(2)

    rss_blocks = scrape_rss()
    all_blocks.extend(rss_blocks)

    return all_blocks


# ---- Offer Parsing ----------------------------------------------------------
def parse_offer(text: str) -> dict:
    lower = text.lower()

    weeks = ""
    m = WEEKS_PAT.search(lower)
    if m:
        weeks = m.group(1)

    waiting = ""
    m_wait = WAITING_PAT.search(lower)
    if m_wait:
        waiting = f"{m_wait.group(1)} & {m_wait.group(2)} month waits waived"
    else:
        m_wait2 = WAITING_PAT2.search(lower)
        if m_wait2:
            waiting = f"{m_wait2.group(1)} month waits waived"

    end_date = ""
    m_date = ENDDATE_PAT.search(text)
    if m_date:
        end_date = m_date.group(1).strip().title()
    else:
        m_date2 = ENDDATE_PAT2.search(text)
        if m_date2:
            end_date = m_date2.group(1).strip().title()
        else:
            m_date3 = ENDDATE_PAT3.search(text)
            if m_date3:
                end_date = m_date3.group(1).strip().title() + " 2026"

    other_parts = []
    m_promo = PROMO_PAT.search(text)
    if m_promo:
        other_parts.append(f"Code: {m_promo.group(1)}")

    return {
        "weeks_free":    weeks,
        "waiting_waive": waiting,
        "other":         " | ".join(other_parts),
        "end_date":      end_date,
    }


def detect_cover_type(text: str) -> str | None:
    lower = text.lower()
    for cover_type, keywords in COVER_TYPE_KEYWORDS.items():
        if any(kw in lower for kw in keywords):
            return cover_type
    return None


def build_structured_offers(blocks: list[str]) -> dict:
    grouped = defaultdict(list)
    seen_parsed = set()
    for block in blocks:
        cover_type = detect_cover_type(block)
        if not cover_type:
            continue
        parsed = parse_offer(block)
        if not any([parsed["weeks_free"], parsed["waiting_waive"],
                    parsed["other"], parsed["end_date"]]):
            continue
        # Skip incomplete fragments
        if not parsed["weeks_free"] and not (parsed["other"] and parsed["end_date"]):
            continue
        # Normalize key
        normalized = re.sub(r'\$[\d,]+', '', parsed["other"])
        normalized = re.sub(r'ends?\s+\d{1,2}\s+\w+\s+\d{4}', '', normalized, flags=re.I)
        normalized = re.sub(r'Code:\s*[A-Z0-9]+', '', normalized, flags=re.I)
        normalized = re.sub(r'\s+', ' ', normalized).strip(" |")
        key = (cover_type, parsed["weeks_free"], normalized)
        if key in seen_parsed:
            continue
        seen_parsed.add(key)
        grouped[cover_type].append(parsed)
    return grouped


def merge_offers(offers: list[dict]) -> dict:
    if not offers:
        return {"weeks_free": "", "waiting_waive": "", "other": "", "end_date": ""}
    if len(offers) == 1:
        return offers[0]

    primary = next((o for o in offers if o["weeks_free"]), offers[0])
    secondary = [o for o in offers if o is not primary]

    merged = {
        "weeks_free":    primary["weeks_free"],
        "waiting_waive": primary["waiting_waive"],
        "other":         primary["other"],
        "end_date":      primary["end_date"],
    }

    seen_parts = set(filter(None, merged["other"].split(" | ")))
    for sec in secondary:
        parts = []
        if sec["other"] and sec["other"] not in seen_parts:
            parts.append(sec["other"])
            seen_parts.add(sec["other"])
        if sec["weeks_free"]:
            wf = f"{sec['weeks_free']} weeks free"
            if wf not in seen_parts:
                parts.append(wf)
                seen_parts.add(wf)
        if sec["end_date"]:
            ed = f"ends {sec['end_date']}"
            if ed not in seen_parts:
                parts.append(ed)
                seen_parts.add(ed)
        if parts:
            merged["other"] = " | ".join(filter(None, [merged["other"]] + parts))

    return merged


# ---- Excel Helpers ----------------------------------------------------------
def _header_index(ws) -> dict:
    idx = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in EXCEL_COL_MAP.values():
                idx[cell.value.strip()] = cell.column - 1
        if len(idx) == len(EXCEL_COL_MAP):
            break
    return idx


def fill_excel(wb: openpyxl.Workbook, structured_offers: dict) -> None:
    ws = wb["table"]
    col_idx = _header_index(ws)

    for row in ws.iter_rows(min_row=2):
        brand_val   = str(row[0].value or "").strip().upper()
        cover_val   = str(row[1].value or "").strip().lower()
        channel_val = str(row[2].value or "").strip().lower()

        if brand_val != "NIB" or channel_val != "direct":
            continue

        # Clear first to avoid stale data
        for field, col_name in EXCEL_COL_MAP.items():
            if col_name in col_idx:
                row[col_idx[col_name]].value = None

        offers_for_cover = structured_offers.get(cover_val, [])
        if not offers_for_cover:
            weeks_col = EXCEL_COL_MAP["weeks_free"]
            if weeks_col in col_idx:
                row[col_idx[weeks_col]].value = "No current offer"
            continue

        merged = merge_offers(offers_for_cover)
        for field, col_name in EXCEL_COL_MAP.items():
            if col_name in col_idx:
                value = merged.get(field)
                row[col_idx[col_name]].value = value if value else None

    print("✓ Filled NIB Direct rows in Excel")


# ---- S3 Helpers -------------------------------------------------------------
def update_on_s3(structured_offers: dict) -> None:
    s3 = boto3.client("s3", region_name="us-east-1")
    key = f"raw/excel/{SHEET}"
    obj = s3.get_object(Bucket=BUCKET, Key=key)
    wb = openpyxl.load_workbook(io.BytesIO(obj["Body"].read()))
    fill_excel(wb, structured_offers)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    s3.put_object(
        Bucket=BUCKET,
        Key=key,
        Body=buf.read(),
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    print(f"✓ Updated: s3://{BUCKET}/{key}")


def update_excel_local(structured_offers: dict, directory: str = ".") -> None:
    src = os.path.join(directory, SHEET)
    if not os.path.exists(src):
        print(f"⚠ {src} not found — skipping Excel update")
        return
    wb = openpyxl.load_workbook(src)
    fill_excel(wb, structured_offers)
    wb.save(src)
    print(f"✓ Updated locally: {src}")


# ---- Output Helpers ---------------------------------------------------------
def build_payload(content: str) -> dict:
    return {
        "source":     SOURCE,
        "tier":       TIER,
        "dataset":    DATASET,
        "scraped_at": datetime.now(timezone.utc).isoformat(),
        "url":        "https://www.nib.com.au/health-insurance/offers",
        "content":    content,
    }


def save_local(payload: dict, directory: str = ".") -> None:
    os.makedirs(directory, exist_ok=True)
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    path = f"{directory}/{SOURCE}_{DATASET}_{date}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"✓ Saved locally: {path}")


def upload_json_to_s3(payload: dict) -> None:
    s3 = boto3.client("s3", region_name="us-east-1")
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    key = f"raw/json/{SOURCE}_{DATASET}_{date}.json"
    s3.put_object(
        Bucket=BUCKET,
        Key=key,
        Body=json.dumps(payload, ensure_ascii=False),
        ContentType="application/json",
    )
    print(f"✓ Uploaded JSON: s3://{BUCKET}/{key}")


# ---- Main -------------------------------------------------------------------
if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--local", metavar="DIR", nargs="?", const="data/competitors",
                        help="save locally instead of uploading to S3")
    args = parser.parse_args()

    blocks = scrape_all_pages()
    structured_offers = build_structured_offers(blocks)

    print(f"\n✓ Structured offers found for: {list(structured_offers.keys())}")

    content = "\n\n".join(blocks)
    payload = build_payload(content)

    if args.local:
        save_local(payload, args.local)
        update_excel_local(structured_offers, args.local)
    else:
        upload_json_to_s3(payload)
        update_on_s3(structured_offers)