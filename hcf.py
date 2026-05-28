# hcf.py — Competitor offer scraper for HCF Health Insurance
# Scrapes direct offer data from hcf.com.au and fills Direct rows in comp_offer.xlsx

import re
import io
import os
import json
import time
import argparse
import feedparser
import requests
import openpyxl
import boto3
from bs4 import BeautifulSoup
from datetime import datetime, timezone
from collections import defaultdict

# ---- Config -----------------------------------------------------------------
SOURCE  = "hcf"
DATASET = "competitor_offers"
TIER    = "phi_industry"
BUCKET  = "p000268ds-comp-offers"
HEADERS = {"User-Agent": "Mozilla/5.0"}
SHEET   = "comp_offer.xlsx"

OFFER_PAGES = [
    {"name": "HCF Homepage",            "url": "https://www.hcf.com.au/"},
    {"name": "HCF Find Health Insurance","url": "https://www.hcf.com.au/health-insurance/find-health-insurance"},
    {"name": "HCF Hospital & Extras",   "url": "https://www.hcf.com.au/health-insurance/find-health-insurance/hospital-and-extras"},
    {"name": "HCF Hospital Only",       "url": "https://www.hcf.com.au/health-insurance/find-health-insurance/hospital"},
    {"name": "HCF Extras Only",         "url": "https://www.hcf.com.au/health-insurance/find-health-insurance/extras"},
]

OFFER_KEYWORDS = ["weeks free", "week free", "waiting period", "waiver", "promo",
                  "offer", "bonus", "discount", "join by", "ends", "new members",
                  "gift card", "mastercard", "e-gift", "cashback", "loyalty"]

TC_KEYWORDS = ["new members only", "eligible members", "eligible customers",
               "fulfilled", "ineligible", "residency", "exclusions apply",
               "waiting periods", "direct debit", "maintained", "australian resident",
               "not have been", "cooling off"]

# ---- Regex Patterns ---------------------------------------------------------
WEEKS_PAT    = re.compile(r'(?:up to\s+)?(\d+(?:\+\d+)?)\s*weeks?\s*free', re.I)
WAITING_PAT  = re.compile(r'(\d+)\s*(?:and|&)\s*(\d+)\s*month.*?(?:wait|waiv)', re.I)
WAITING_PAT2 = re.compile(r'(\d+)\s*month.*?(?:wait|waiv)', re.I)
ENDDATE_PAT  = re.compile(r'(?:ends?|until|by|join by|closes?)\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4}|\d{1,2}\s+[A-Za-z]+\s*\d{2,4})', re.I)
ENDDATE_PAT2 = re.compile(r'(\d{1,2}\s+[A-Za-z]{3,}\s+\d{4})', re.I)
ENDDATE_PAT3 = re.compile(r'(?:ends?|until|by|join by|closes?)\s+(\d{1,2}\s+[A-Za-z]+)(?:\s|\.)', re.I)
GIFT_PAT     = re.compile(r'\$[\d,]+\s*(?:digital\s*prepaid\s*mastercard|gift\s*card|e-gift|visa|mastercard|cashback)', re.I)
PROMO_PAT    = re.compile(r'(?:promo(?:tion)?\s*code|use\s+(?:promo\s+)?code|enter\s+code|code)[:\s]+([A-Z0-9]{6,})', re.I)

EXCEL_COL_MAP = {
    "weeks_free":    "Offer : Weeks Free",
    "waiting_waive": "Offer : Waiting period waive",
    "other":         "Offer : Other",
    "end_date":      "Offer : End date",
    "tandc":         "Offer : T&C",
}

COVER_TYPE_KEYWORDS = {
    "hospital + extras": ["hospital and extras", "hospital & extras", "combined",
                          "hospital + extras", "hospital and extra"],
    "hospital only":     ["hospital only", "hospital cover", "hospital-only",
                          "standalone hospital", "hospital product"],
    "extras only":       ["extras only", "extras cover", "extras-only",
                          "standalone extras", "extras product"],
}

# ---- Scraping ---------------------------------------------------------------
def fetch_page(url: str) -> BeautifulSoup | None:
    try:
        r = requests.get(url, headers=HEADERS, timeout=30)
        if r.status_code == 200:
            return BeautifulSoup(r.text, "html.parser")
        print(f"  ✗ {url}: {r.status_code}")
    except Exception as e:
        print(f"  ✗ {url}: {e}")
    return None


def extract_offer_blocks(soup: BeautifulSoup) -> list[str]:
    found = []
    for tag in soup.find_all(["p", "h1", "h2", "h3", "h4", "li", "span", "div"]):
        text = tag.get_text(separator=" ", strip=True)
        if any(kw in text.lower() for kw in OFFER_KEYWORDS):
            if 20 < len(text) < 600 and text not in found:
                found.append(text)
    return found[:30]


def scrape_rss() -> list[str]:
    print("  Fetching HCF offers via Google News RSS...")
    rss_url = "https://news.google.com/rss/search?q=HCF+health+insurance+offer+weeks+free+Australia&hl=en-AU&gl=AU&ceid=AU:en"
    feed = feedparser.parse(rss_url)
    entries = []
    for entry in feed.entries[:10]:
        entries.append(f"{entry.title} | {entry.link} | {entry.get('published', '')}")
    print(f"  ✓ Found {len(entries)} RSS articles")
    return entries


def scrape_all_pages() -> tuple[list[str], dict]:
    print("--- HCF Scraper ---")
    all_blocks = []
    tandc_by_cover = defaultdict(list)

    PAGE_COVER_MAP = {
        "HCF Homepage":             ["hospital + extras", "hospital only", "extras only"],
        "HCF Find Health Insurance":["hospital + extras", "hospital only", "extras only"],
        "HCF Hospital & Extras":    ["hospital + extras"],
        "HCF Hospital Only":        ["hospital only"],
        "HCF Extras Only":          ["extras only"],
    }

    for page in OFFER_PAGES:
        print(f"  Fetching {page['name']}...")
        soup = fetch_page(page["url"])
        if soup:
            blocks = extract_offer_blocks(soup)
            all_blocks.extend(blocks)
            for ct in PAGE_COVER_MAP.get(page["name"], []):
                text = extract_tandc_text(soup, ct)
                if text:
                    tandc_by_cover[ct].append(text)
            print(f"  ✓ {page['name']}: {len(blocks)} offer blocks found")
        else:
            print(f"  ✗ {page['name']}: failed")
        time.sleep(2)

    rss_blocks = scrape_rss()
    all_blocks.extend(rss_blocks)

    tandc_final = {ct: texts[0] for ct, texts in tandc_by_cover.items() if texts}

    return all_blocks, tandc_final


# ---- T&C Extraction ---------------------------------------------------------
def parse_tandc_structured(text: str, cover_type: str = "") -> str:
    lower = text.lower()
    parts = []
    seen_labels = set()

    def add(label, value):
        if label not in seen_labels and value:
            seen_labels.add(label)
            parts.append(f"{label}: {value.strip().title()}")

    # Eligibility
    elig_match = re.search(r'(new members? only|eligible members?|eligible customers?|australian residents?)', lower)
    if elig_match:
        add("Eligibility", elig_match.group(1))

    # Offer period
    period_match = re.search(r'(?:join|take out)(?:\s+eligible)?\s+(?:combined\s+)?(?:hospital[^.]*)?(?:cover\s+)?by\s+(\d{1,2}\s+[A-Za-z]+(?:\s+\d{4})?)', lower)
    if period_match:
        date = period_match.group(1).strip()
        if not re.search(r'\d{4}', date):
            date += " 2026"
        add("Offer Period", f"Join by {date}")

    # Waiting periods
    if cover_type != "hospital only":
        wait_match = re.search(r'(waive?\s*the?\s*\d+[-\s]month[^.]*|skip\s*the?\s*\d+[^.]*month[^.]*waiting[^.]*)', lower)
        if wait_match:
            add("Waiting Periods", wait_match.group(1))

    # Fulfilment
    fulfil_match = re.search(r'(?:after|following)\s+(\d+\s*(?:days?|months?|years?))[^.]*(?:free|weeks)', lower)
    if fulfil_match:
        add("Fulfilment", f"After {fulfil_match.group(1)}")

    # Gift card / reward
    gift_match = re.search(r'(\$[\d,]+\s*(?:digital\s*prepaid\s*mastercard|gift\s*card|e-gift|mastercard|cashback))', lower)
    if gift_match:
        add("Reward", gift_match.group(1).title())

    # Excluded covers
    excl_match = re.search(r'exclu(?:ding|des?|sions?)[:\s]+([^.]{10,})', lower)
    if excl_match:
        add("Excluded", excl_match.group(1))

    return " | ".join(parts) if parts else text[:200]


def extract_tandc_text(soup: BeautifulSoup, cover_type: str) -> str:
    found = []
    for tag in soup.find_all(["p", "li", "span"]):
        text = tag.get_text(separator=" ", strip=True)
        text = re.sub(r',?\s*opens?\s+in\s+a\s+new\s+tab', '', text, flags=re.I).strip()
        text = re.sub(r'T&Cs?\s*(?:apply\.?)?', '', text, flags=re.I).strip()
        text = re.sub(r'\s+', ' ', text).strip()
        lower = text.lower()
        if not any(kw in lower for kw in TC_KEYWORDS):
            continue
        if not (20 < len(text) < 300):
            continue
        if text in found:
            continue
        if cover_type == "hospital only":
            if any(k in lower for k in ["hospital only", "hospital-only"]) and "extras" not in lower:
                found.append(text)
            elif any(k in lower for k in ["eligible members", "new members"]) and "extras" not in lower:
                found.append(text)
        elif cover_type == "hospital + extras":
            if any(k in lower for k in ["combined", "hospital and extras", "hospital + extras",
                                          "hospital & extras", "hospital and extras"]):
                found.append(text)
            elif any(k in lower for k in ["eligible members", "new members", "waiting period"]):
                found.append(text)
        elif cover_type == "extras only":
            if any(k in lower for k in ["extras only", "extras cover"]):
                found.append(text)

    if not found:
        return ""

    combined = " ".join(list(dict.fromkeys(found)))
    return parse_tandc_structured(combined, cover_type)


# ---- Offer Parsing ----------------------------------------------------------
def parse_offer(text: str) -> dict:
    lower = text.lower()

    weeks = ""
    m = WEEKS_PAT.search(lower)
    if m:
        weeks = m.group(1)

    waiting = ""
    m_wait = WAITING_PAT.search(lower)
    if m_wait:
        waiting = f"{m_wait.group(1)} & {m_wait.group(2)} month waits waived"
    else:
        m_wait2 = WAITING_PAT2.search(lower)
        if m_wait2:
            waiting = f"{m_wait2.group(1)} month waits waived"

    end_date = ""
    m_date = ENDDATE_PAT.search(lower)
    if m_date:
        end_date = m_date.group(1).strip().title()
    else:
        m_date2 = ENDDATE_PAT2.search(lower)
        if m_date2:
            end_date = m_date2.group(1).strip().title()
        else:
            m_date3 = ENDDATE_PAT3.search(lower)
            if m_date3:
                end_date = m_date3.group(1).strip().title() + " 2026"

    other_parts = []
    m_gift = GIFT_PAT.search(text)
    if m_gift:
        other_parts.append(m_gift.group(0).strip())
    m_promo = PROMO_PAT.search(text)
    if m_promo:
        other_parts.append(f"Code: {m_promo.group(1)}")

    return {
        "weeks_free":    weeks,
        "waiting_waive": waiting,
        "other":         " | ".join(other_parts),
        "end_date":      end_date,
        "tandc":         "",
    }


def detect_cover_type(text: str) -> str | None:
    lower = text.lower()
    for cover_type, keywords in COVER_TYPE_KEYWORDS.items():
        if any(kw in lower for kw in keywords):
            return cover_type
    return None


def build_structured_offers(blocks: list[str], tandc_by_cover: dict = {}) -> dict:
    grouped = defaultdict(list)
    seen_parsed = set()
    for block in blocks:
        cover_type = detect_cover_type(block)
        if not cover_type:
            continue
        parsed = parse_offer(block)
        if not any([parsed["weeks_free"], parsed["waiting_waive"],
                    parsed["other"], parsed["end_date"]]):
            continue
        if not parsed["weeks_free"] and not (parsed["other"] and parsed["end_date"]):
            continue
        normalized = re.sub(r'\$[\d,]+', '', parsed["other"])
        normalized = re.sub(r'ends?\s+\d{1,2}\s+\w+\s+\d{4}', '', normalized, flags=re.I)
        normalized = re.sub(r'Code:\s*[A-Z0-9]+', '', normalized, flags=re.I)
        normalized = re.sub(r'\s+', ' ', normalized).strip(" |")
        key = (cover_type, parsed["weeks_free"], normalized)
        if key in seen_parsed:
            continue
        seen_parsed.add(key)
        tandc = tandc_by_cover.get(cover_type, "")
        if tandc and "Offer Period" not in tandc and parsed["end_date"]:
            tandc = tandc + f" | Offer Period: Join By {parsed['end_date']}"
        elif not tandc and parsed["end_date"]:
            tandc = f"Offer Period: Join By {parsed['end_date']}"
        parsed["tandc"] = tandc
        grouped[cover_type].append(parsed)

    return grouped


def merge_offers(offers: list[dict]) -> dict:
    seen = set()
    unique_offers = []
    for offer in offers:
        key = (offer["weeks_free"], offer["other"], offer["end_date"])
        if key not in seen and any([offer["weeks_free"], offer["waiting_waive"],
                                    offer["other"], offer["end_date"]]):
            seen.add(key)
            unique_offers.append(offer)

    if not unique_offers:
        return {"weeks_free": "", "waiting_waive": "", "other": "", "end_date": "", "tandc": ""}
    if len(unique_offers) == 1:
        return unique_offers[0]

    primary = next((o for o in unique_offers if o["weeks_free"]), unique_offers[0])
    secondary = [o for o in unique_offers if o is not primary]

    merged = {
        "weeks_free":    primary["weeks_free"],
        "waiting_waive": primary["waiting_waive"],
        "other":         primary["other"],
        "end_date":      primary["end_date"],
        "tandc":         primary.get("tandc", ""),
    }

    seen_parts = set(filter(None, merged["other"].split(" | ")))
    for sec in secondary:
        parts = []
        sec_other_clean = re.sub(r',?\s*ends\s+\d{1,2}\s+\w+\s+\d{4}', '', sec.get("other", ""), flags=re.I).strip(" ,|")
        if sec_other_clean and not sec_other_clean.rstrip().endswith(':') and sec_other_clean not in seen_parts:
            parts.append(sec_other_clean)
            seen_parts.add(sec_other_clean)
        if sec["weeks_free"]:
            wf = f"{sec['weeks_free']} weeks free"
            if wf not in seen_parts:
                parts.append(wf)
                seen_parts.add(wf)
        if sec["end_date"] and sec["end_date"] not in merged["end_date"]:
            merged["end_date"] = f"{merged['end_date']} | Alt: {sec['end_date']}"
        if parts:
            merged["other"] = " | ".join(filter(None, [merged["other"]] + parts))

    return merged


# ---- Excel Helpers ----------------------------------------------------------
def _header_index(ws) -> dict:
    idx = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in EXCEL_COL_MAP.values():
                idx[cell.value.strip()] = cell.column - 1
        if len(idx) == len(EXCEL_COL_MAP):
            break
    return idx


def fill_excel(wb: openpyxl.Workbook, structured_offers: dict) -> None:
    ws = wb["table"]
    col_idx = _header_index(ws)

    for row in ws.iter_rows(min_row=2):
        brand_val   = str(row[0].value or "").strip().upper()
        cover_val   = str(row[1].value or "").strip().lower()
        channel_val = str(row[2].value or "").strip().lower()

        if brand_val != "HCF" or channel_val != "direct":
            continue

        for field, col_name in EXCEL_COL_MAP.items():
            if col_name in col_idx:
                row[col_idx[col_name]].value = None

        offers_for_cover = structured_offers.get(cover_val, [])
        if not offers_for_cover:
            weeks_col = EXCEL_COL_MAP["weeks_free"]
            if weeks_col in col_idx:
                row[col_idx[weeks_col]].value = "No current offer"
            continue

        merged = merge_offers(offers_for_cover)
        for field, col_name in EXCEL_COL_MAP.items():
            if col_name in col_idx:
                value = merged.get(field)
                row[col_idx[col_name]].value = value if value else None

    print("✓ Filled HCF Direct rows in Excel")


# ---- S3 Helpers -------------------------------------------------------------
def update_on_s3(structured_offers: dict) -> None:
    s3 = boto3.client("s3", region_name="us-east-1")
    key = f"raw/excel/{SHEET}"
    obj = s3.get_object(Bucket=BUCKET, Key=key)
    wb = openpyxl.load_workbook(io.BytesIO(obj["Body"].read()))
    fill_excel(wb, structured_offers)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    s3.put_object(
        Bucket=BUCKET,
        Key=key,
        Body=buf.read(),
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    print(f"✓ Updated: s3://{BUCKET}/{key}")


def update_excel_local(structured_offers: dict, directory: str = ".") -> None:
    src = os.path.join(directory, SHEET)
    if not os.path.exists(src):
        print(f"⚠ {src} not found — skipping Excel update")
        return
    wb = openpyxl.load_workbook(src)
    fill_excel(wb, structured_offers)
    wb.save(src)
    print(f"✓ Updated locally: {src}")


# ---- Output Helpers ---------------------------------------------------------
def build_payload(content: str) -> dict:
    return {
        "source":     SOURCE,
        "tier":       TIER,
        "dataset":    DATASET,
        "scraped_at": datetime.now(timezone.utc).isoformat(),
        "url":        "https://www.hcf.com.au/health-insurance",
        "content":    content,
    }


def save_local(payload: dict, directory: str = ".") -> None:
    os.makedirs(directory, exist_ok=True)
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    path = f"{directory}/{SOURCE}_{DATASET}_{date}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"✓ Saved locally: {path}")


def upload_json_to_s3(payload: dict) -> None:
    s3 = boto3.client("s3", region_name="us-east-1")
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    key = f"raw/json/{SOURCE}_{DATASET}_{date}.json"
    s3.put_object(
        Bucket=BUCKET,
        Key=key,
        Body=json.dumps(payload, ensure_ascii=False),
        ContentType="application/json",
    )
    print(f"✓ Uploaded JSON: s3://{BUCKET}/{key}")


# ---- Main -------------------------------------------------------------------
if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--local", metavar="DIR", nargs="?", const="data/competitors",
                        help="save locally instead of uploading to S3")
    args = parser.parse_args()

    blocks, tandc_by_cover = scrape_all_pages()
    print(f"\n✓ T&C text found for: {list(tandc_by_cover.keys())}")
    structured_offers = build_structured_offers(blocks, tandc_by_cover)
    print(f"\n✓ Structured offers found for: {list(structured_offers.keys())}")

    content = "\n\n".join(blocks)
    payload = build_payload(content)

    if args.local:
        save_local(payload, args.local)
        update_excel_local(structured_offers, args.local)
    else:
        upload_json_to_s3(payload)
        update_on_s3(structured_offers)